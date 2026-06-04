"""Ingest files from input/ folder: move md/txt to raw/, convert docx/pdf to clean markdown."""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

import fitz  # pymupdf
from docx import Document
from pydantic import BaseModel
from rich.console import Console

console = Console()


class IngestConfig(BaseModel):
    input_dir: Path = Path("input")
    raw_dir: Path = Path("raw")


SUPPORTED_DIRECT = {".md", ".txt"}
SUPPORTED_CONVERT = {".docx", ".pdf"}


def docx_to_markdown(filepath: Path) -> str:
    doc = Document(str(filepath))
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue

        style = para.style.name.lower() if para.style else ""
        if "heading 1" in style:
            lines.append(f"# {text}")
        elif "heading 2" in style:
            lines.append(f"## {text}")
        elif "heading 3" in style:
            lines.append(f"### {text}")
        elif "heading 4" in style:
            lines.append(f"#### {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)

    # Process tables
    for table in doc.tables:
        lines.append("")
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if i == 0:
                lines.append("| " + " | ".join("---" for _ in cells) + " |")
        lines.append("")

    md = "\n".join(lines)
    # Clean up Word artifacts
    md = re.sub(r"\n{3,}", "\n\n", md)
    md = re.sub(r"<[^>]+>", "", md)  # strip any XML/HTML tags
    md = md.strip()
    return md


def pdf_to_markdown(filepath: Path) -> str:
    doc = fitz.open(str(filepath))
    pages: list[str] = []

    # Detect repeated headers/footers by finding text blocks that appear on 3+ pages
    all_blocks: list[list[str]] = []
    for page in doc:
        blocks = [b[4].strip() for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
        all_blocks.append(blocks)

    # Find repeated lines (headers/footers)
    repeated: set[str] = set()
    if len(all_blocks) >= 3:
        from collections import Counter
        line_counts: Counter[str] = Counter()
        for blocks in all_blocks:
            # Check first 2 and last 2 blocks per page (likely header/footer)
            candidates = blocks[:2] + blocks[-2:] if len(blocks) > 4 else blocks
            for line in candidates:
                if len(line) < 200:  # headers/footers are short
                    line_counts[line] += 1
        for line, count in line_counts.items():
            if count >= min(3, len(all_blocks)):
                repeated.add(line)

    # Extract text, skipping repeated headers/footers
    for blocks in all_blocks:
        page_text: list[str] = []
        for block in blocks:
            if block not in repeated:
                page_text.append(block)
        if page_text:
            pages.append("\n\n".join(page_text))

    doc.close()

    md = "\n\n---\n\n".join(pages)
    md = re.sub(r"\n{3,}", "\n\n", md)
    md = md.strip()
    return md


def ingest_file(filepath: Path, raw_dir: Path) -> tuple[str, str]:
    suffix = filepath.suffix.lower()
    stem = filepath.stem
    target = raw_dir / f"{stem}.md"

    if suffix in SUPPORTED_DIRECT:
        # Move directly (rename .txt to .md)
        shutil.move(str(filepath), str(target))
        return (filepath.name, "moved")

    if suffix == ".docx":
        md = docx_to_markdown(filepath)
        target.write_text(md, encoding="utf-8")
        filepath.unlink()
        return (filepath.name, "converted from docx")

    if suffix == ".pdf":
        md = pdf_to_markdown(filepath)
        target.write_text(md, encoding="utf-8")
        filepath.unlink()
        return (filepath.name, "converted from pdf")

    return (filepath.name, f"skipped (unsupported: {suffix})")


def ingest(config: IngestConfig) -> None:
    config.input_dir.mkdir(parents=True, exist_ok=True)
    config.raw_dir.mkdir(parents=True, exist_ok=True)

    files = [f for f in config.input_dir.iterdir() if f.is_file()]
    if not files:
        console.print("[dim]No files in input/ folder.[/dim]")
        return

    console.print(f"[cyan]Found {len(files)} files in {config.input_dir}[/cyan]")

    moved = 0
    converted = 0
    skipped = 0

    for f in sorted(files):
        name, action = ingest_file(f, config.raw_dir)
        if "moved" in action:
            console.print(f"  [green]{name}[/green] → {action}")
            moved += 1
        elif "converted" in action:
            console.print(f"  [green]{name}[/green] → {action}")
            converted += 1
        else:
            console.print(f"  [yellow]{name}[/yellow] → {action}")
            skipped += 1

    console.print()
    console.print(f"[bold green]Done![/bold green] {moved} moved, {converted} converted, {skipped} skipped")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Ingest files from input/ folder into raw/ as clean markdown."
    )
    parser.add_argument("--input-dir", type=Path, default="input", help="Input directory (default: input)")
    parser.add_argument("--raw-dir", type=Path, default="raw", help="Raw output directory (default: raw)")

    args = parser.parse_args()
    config = IngestConfig(input_dir=args.input_dir, raw_dir=args.raw_dir)
    ingest(config)


if __name__ == "__main__":
    main()
