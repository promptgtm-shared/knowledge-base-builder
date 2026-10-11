"""Content conversion: HTML pages and local documents to clean markdown (no links)."""

from __future__ import annotations

import re
import warnings
from collections import Counter
from pathlib import Path

import html2text
import pymupdf
from bs4 import BeautifulSoup, XMLParsedAsHTMLWarning
from docx import Document

warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)

CONTENT_SELECTORS = [
    ".entry-content",
    "article",
    "main",
    "[role='main']",
    ".post-content",
    ".page-content",
    ".content",
]


def extract_content(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")

    for tag in soup.find_all(["script", "style", "nav", "footer", "header", "aside", "noscript"]):
        tag.decompose()

    for selector in CONTENT_SELECTORS:
        element = soup.select_one(selector)
        if element and len(element.get_text(strip=True)) > 100:
            return str(element)

    body = soup.find("body")
    return str(body) if body else html


def html_to_markdown(html: str) -> str:
    content_html = extract_content(html)

    h = html2text.HTML2Text()
    h.ignore_links = True
    h.ignore_images = True
    h.ignore_emphasis = False
    h.body_width = 0
    h.unicode_snob = True
    h.skip_internal_links = True
    h.single_line_break = False
    h.protect_links = False
    h.wrap_links = False
    h.wrap_list_items = False

    md = h.handle(content_html)

    md = re.sub(r"\n{3,}", "\n\n", md)
    md = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", md)
    md = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", md)
    return md.strip()


def docx_to_markdown(filepath: Path) -> str:
    doc = Document(str(filepath))
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue

        style = para.style.name.lower() if para.style else ""
        if "heading 1" in style:
            lines.append(f"# {text}")
        elif "heading 2" in style:
            lines.append(f"## {text}")
        elif "heading 3" in style:
            lines.append(f"### {text}")
        elif "heading 4" in style:
            lines.append(f"#### {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)

    for table in doc.tables:
        lines.append("")
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if i == 0:
                lines.append("| " + " | ".join("---" for _ in cells) + " |")
        lines.append("")

    md = "\n".join(lines)
    md = re.sub(r"\n{3,}", "\n\n", md)
    md = re.sub(r"<[^>]+>", "", md)
    return md.strip()


def pdf_to_markdown(filepath: Path) -> str:
    doc = pymupdf.open(str(filepath))

    all_blocks: list[list[str]] = []
    for page in doc:
        blocks = [b[4].strip() for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
        all_blocks.append(blocks)

    # Repeated short lines near page edges are headers/footers
    repeated: set[str] = set()
    if len(all_blocks) >= 3:
        line_counts: Counter[str] = Counter()
        for blocks in all_blocks:
            candidates = blocks[:2] + blocks[-2:] if len(blocks) > 4 else blocks
            for line in candidates:
                if len(line) < 200:
                    line_counts[line] += 1
        for line, count in line_counts.items():
            if count >= min(3, len(all_blocks)):
                repeated.add(line)

    pages: list[str] = []
    for blocks in all_blocks:
        page_text = [block for block in blocks if block not in repeated]
        if page_text:
            pages.append("\n\n".join(page_text))

    doc.close()

    md = "\n\n---\n\n".join(pages)
    md = re.sub(r"\n{3,}", "\n\n", md)
    return md.strip()


DOCUMENT_CONVERTERS = {".docx": docx_to_markdown, ".pdf": pdf_to_markdown}
TEXT_SUFFIXES = {".md", ".txt"}


def document_to_markdown(filepath: Path) -> str | None:
    """Convert a local document to markdown; None for unsupported types."""
    suffix = filepath.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return filepath.read_text(encoding="utf-8").strip()
    converter = DOCUMENT_CONVERTERS.get(suffix)
    return converter(filepath) if converter else None
